"""
Service d'export pour générer des fichiers PDF et DOCX
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

# Pour PDF
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

# Pour DOCX
try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

logger = logging.getLogger(__name__)


class ExportService:
    """
    Service pour exporter les romans en PDF et DOCX
    """

    def __init__(self):
        self.check_dependencies()
        logger.info("ExportService initialisé")

    def check_dependencies(self):
        """Vérifie que les dépendances sont installées"""
        if not REPORTLAB_AVAILABLE:
            logger.warning("reportlab non installé - export PDF désactivé")
        if not PYTHON_DOCX_AVAILABLE:
            logger.warning("python-docx non installé - export DOCX désactivé")

    def export_to_pdf(
        self,
        novel,
        output_path: str,
        include_cover: bool = True,
        include_toc: bool = True,
        author_name: str = "",
        subtitle: str = ""
    ) -> bool:
        """
        Exporte le roman en PDF

        Args:
            novel: Roman à exporter
            output_path: Chemin du fichier de sortie
            include_cover: Inclure page de couverture
            include_toc: Inclure table des matières
            author_name: Nom de l'auteur
            subtitle: Sous-titre du roman

        Returns:
            True si succès, False sinon
        """
        if not REPORTLAB_AVAILABLE:
            logger.error("reportlab non installé")
            return False

        try:
            # Créer le document PDF
            doc = SimpleDocTemplate(
                output_path,
                pagesize=A4,
                rightMargin=2*cm,
                leftMargin=2*cm,
                topMargin=2*cm,
                bottomMargin=2*cm
            )

            # Styles
            styles = getSampleStyleSheet()

            # Style pour le titre principal
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor='#2c3e50',
                spaceAfter=30,
                alignment=TA_CENTER,
                fontName='Helvetica-Bold'
            )

            # Style pour le sous-titre
            subtitle_style = ParagraphStyle(
                'CustomSubtitle',
                parent=styles['Normal'],
                fontSize=14,
                textColor='#7f8c8d',
                spaceAfter=12,
                alignment=TA_CENTER,
                fontName='Helvetica-Oblique'
            )

            # Style pour l'auteur
            author_style = ParagraphStyle(
                'CustomAuthor',
                parent=styles['Normal'],
                fontSize=16,
                textColor='#34495e',
                spaceAfter=40,
                alignment=TA_CENTER,
                fontName='Helvetica'
            )

            # Style pour les titres de chapitre
            chapter_style = ParagraphStyle(
                'CustomChapter',
                parent=styles['Heading2'],
                fontSize=18,
                textColor='#2980b9',
                spaceAfter=20,
                spaceBefore=30,
                fontName='Helvetica-Bold'
            )

            # Style pour le corps du texte
            body_style = ParagraphStyle(
                'CustomBody',
                parent=styles['Normal'],
                fontSize=12,
                leading=18,
                alignment=TA_JUSTIFY,
                spaceAfter=12,
                fontName='Helvetica'
            )

            # Construire le contenu
            story = []

            # Page de couverture
            if include_cover:
                story.append(Spacer(1, 5*cm))
                story.append(Paragraph(novel.title, title_style))

                if subtitle:
                    story.append(Paragraph(subtitle, subtitle_style))
                    story.append(Spacer(1, 1*cm))

                if author_name:
                    story.append(Spacer(1, 2*cm))
                    story.append(Paragraph(f"par {author_name}", author_style))

                story.append(Spacer(1, 3*cm))
                story.append(Paragraph(
                    f"Généré le {datetime.now().strftime('%d/%m/%Y')}",
                    styles['Normal']
                ))
                story.append(PageBreak())

            # Table des matières
            if include_toc and novel.chapters:
                story.append(Paragraph("Table des matières", title_style))
                story.append(Spacer(1, 1*cm))

                for idx, chapter in enumerate(novel.chapters, 1):
                    toc_entry = f"{idx}. {chapter.title}"
                    story.append(Paragraph(toc_entry, styles['Normal']))
                    story.append(Spacer(1, 0.3*cm))

                story.append(PageBreak())

            # Chapitres
            for chapter in novel.chapters:
                # Titre du chapitre
                story.append(Paragraph(chapter.title, chapter_style))
                story.append(Spacer(1, 0.5*cm))

                # Contenu du chapitre
                # Traiter le contenu HTML/Rich text
                content = self._clean_html_for_pdf(chapter.content)

                # Diviser en paragraphes
                paragraphs = content.split('\n')
                for para in paragraphs:
                    if para.strip():
                        story.append(Paragraph(para, body_style))

                story.append(PageBreak())

            # Générer le PDF
            doc.build(story)

            logger.info(f"PDF exporté avec succès : {output_path}")
            return True

        except Exception as e:
            logger.error(f"Erreur lors de l'export PDF : {e}")
            return False

    def export_to_docx(
        self,
        novel,
        output_path: str,
        include_cover: bool = True,
        include_toc: bool = True,
        author_name: str = "",
        subtitle: str = ""
    ) -> bool:
        """
        Exporte le roman en DOCX

        Args:
            novel: Roman à exporter
            output_path: Chemin du fichier de sortie
            include_cover: Inclure page de couverture
            include_toc: Inclure table des matières
            author_name: Nom de l'auteur
            subtitle: Sous-titre du roman

        Returns:
            True si succès, False sinon
        """
        if not PYTHON_DOCX_AVAILABLE:
            logger.error("python-docx non installé")
            return False

        try:
            # Créer le document
            doc = Document()

            # Configuration des marges
            sections = doc.sections
            for section in sections:
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1)
                section.right_margin = Inches(1)

            # Page de couverture
            if include_cover:
                # Titre
                title = doc.add_heading(novel.title, level=1)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_run = title.runs[0]
                title_run.font.size = Pt(28)
                title_run.font.color.rgb = RGBColor(44, 62, 80)

                # Sous-titre
                if subtitle:
                    subtitle_para = doc.add_paragraph(subtitle)
                    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    subtitle_run = subtitle_para.runs[0]
                    subtitle_run.font.size = Pt(16)
                    subtitle_run.font.italic = True
                    subtitle_run.font.color.rgb = RGBColor(127, 140, 141)

                # Espacement
                doc.add_paragraph()
                doc.add_paragraph()

                # Auteur
                if author_name:
                    author_para = doc.add_paragraph(f"par {author_name}")
                    author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    author_run = author_para.runs[0]
                    author_run.font.size = Pt(18)
                    author_run.font.color.rgb = RGBColor(52, 73, 94)

                # Date
                doc.add_paragraph()
                doc.add_paragraph()
                date_para = doc.add_paragraph(
                    f"Généré le {datetime.now().strftime('%d/%m/%Y')}"
                )
                date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

                # Saut de page
                doc.add_page_break()

            # Table des matières
            if include_toc and novel.chapters:
                toc_title = doc.add_heading("Table des matières", level=1)
                toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

                doc.add_paragraph()

                for idx, chapter in enumerate(novel.chapters, 1):
                    toc_entry = doc.add_paragraph(f"{idx}. {chapter.title}")
                    toc_entry_run = toc_entry.runs[0]
                    toc_entry_run.font.size = Pt(12)

                doc.add_page_break()

            # Chapitres
            for chapter in novel.chapters:
                # Titre du chapitre
                chapter_heading = doc.add_heading(chapter.title, level=2)
                chapter_run = chapter_heading.runs[0]
                chapter_run.font.size = Pt(20)
                chapter_run.font.color.rgb = RGBColor(41, 128, 185)

                doc.add_paragraph()

                # Contenu du chapitre
                content = self._clean_html_for_docx(chapter.content)

                # Diviser en paragraphes
                paragraphs = content.split('\n')
                for para in paragraphs:
                    if para.strip():
                        p = doc.add_paragraph(para)
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

                        # Formatter le texte
                        for run in p.runs:
                            run.font.size = Pt(12)
                            run.font.name = 'Calibri'

                # Saut de page après chaque chapitre
                doc.add_page_break()

            # Sauvegarder le document
            doc.save(output_path)

            logger.info(f"DOCX exporté avec succès : {output_path}")
            return True

        except Exception as e:
            logger.error(f"Erreur lors de l'export DOCX : {e}")
            return False

    def _clean_html_for_pdf(self, html_content: str) -> str:
        """Nettoie le contenu HTML pour l'export PDF"""
        # Enlever les balises HTML basiques
        import re

        # Remplacer <br> par des sauts de ligne
        content = re.sub(r'<br\s*/?>', '\n', html_content)

        # Enlever les balises HTML restantes (garder le texte)
        content = re.sub(r'<[^>]+>', '', content)

        # Décoder les entités HTML
        content = content.replace('&nbsp;', ' ')
        content = content.replace('&lt;', '<')
        content = content.replace('&gt;', '>')
        content = content.replace('&amp;', '&')
        content = content.replace('&quot;', '"')

        return content.strip()

    def _clean_html_for_docx(self, html_content: str) -> str:
        """Nettoie le contenu HTML pour l'export DOCX"""
        # Même logique que pour PDF
        return self._clean_html_for_pdf(html_content)

    def is_pdf_available(self) -> bool:
        """Retourne True si l'export PDF est disponible"""
        return REPORTLAB_AVAILABLE

    def is_docx_available(self) -> bool:
        """Retourne True si l'export DOCX est disponible"""
        return PYTHON_DOCX_AVAILABLE
